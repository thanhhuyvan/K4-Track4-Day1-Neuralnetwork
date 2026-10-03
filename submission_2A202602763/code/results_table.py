import json
from pathlib import Path
from copy import copy
import numpy as np
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.workbook.properties import CalcProperties

def save_result(result, results_dir="../results"):
    directory = Path(results_dir)
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{result['cfg']['exp_id']}.json"
    payload = {k: v for k, v in result.items() if k != "best_state"}
    path.write_text(
        json.dumps(payload, indent=2, allow_nan=False), encoding="utf-8"
    )
    return str(path)

def load_results(results_dir="../results"):
    results = []
    for path in sorted(Path(results_dir).glob("*.json")):
        result = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(result, dict) and "cfg" in result and "history" in result:
            results.append(result)
    return results

def to_row(result, eval_scores=None, notes=""):
    cfg = result["cfg"]
    history = result["history"]
    last = history[-1] if history else {}
    row = dict(cfg)
    row.update(
        hidden="→".join(map(str, cfg["hidden"])),
        step0_loss=result["initial_val_loss"],
        best_val_loss=result["best_val_loss"],
        best_epoch=result["best_epoch"],
        final_train_loss=last.get("train_loss"),
        final_val_loss=last.get("val_loss"),
        val_acc=result["val_acc"],
        val_macro_f1=result["val_macro_f1"],
        time_per_epoch_s=(
            float(np.mean([r["seconds"] for r in history]))
            if history else None
        ),
        peak_mem_MB=result["peak_memory_mb"],
        diverged=result["diverged"],
        eval_acc=None if eval_scores is None else eval_scores["accuracy"],
        eval_macro_f1=None if eval_scores is None else eval_scores["macro_f1"],
        figure_file=f"figures/{cfg['exp_id']}.png",
        notes=notes,
    )
    if result["diverged"]:
        row["notes"] += " Diverged; metrics may describe an earlier checkpoint."
    return row

def write_xlsx(rows, template_path, out_path):
    wb = load_workbook(template_path)
    ws = wb["Experiments"]
    headers = {cell.column: cell.value for cell in ws[1]}
    formula_columns = {
        "step0_gap_vs_lnC", "gap_val_minus_train",
        "delta_val_f1_vs_base", "beyond_noise",
    }

    # Clear template example values while preserving formulas.
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            if cell.data_type != "f":
                cell.value = None

    for index, row in enumerate(rows, start=2):
        for column, name in headers.items():
            target = ws.cell(index, column)
            source = ws.cell(2, column)
            if index > 2:
                target._style = copy(source._style)
            if name in formula_columns:
                if source.data_type == "f":
                    target.value = Translator(
                        source.value, origin=source.coordinate
                    ).translate_formula(target.coordinate)
            elif name in row:
                target.value = row[name]

    # Template calculations cover rows 2–61.
    assert len(rows) <= 60, "Extend template formula ranges for more than 60 runs"
    wb.calculation = CalcProperties(
        calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True
    )
    wb.save(out_path)
